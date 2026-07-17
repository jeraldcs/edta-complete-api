"""
Generate EDTA Complete Showcase Word document.

Output: docs/EDTA_COMPLETE_SHOWCASE.docx

Combines architecture, demo walkthrough, code map, API integration,
benefits, multi-industry use, roadmap, and publication materials.

Usage:
    python docs/generate_showcase_docx.py
"""

from __future__ import annotations

import subprocess
import sys
from pathlib import Path

import fitz
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / "_showcase_docx_assets"
OUTPUT = ROOT / "EDTA_COMPLETE_SHOWCASE.docx"
FULL_PDF = ROOT / "edta_architecture_diagrams_full.pdf"
LEGACY_PDF = ROOT / "edta_architecture_diagram.pdf"

LIVE_DEMO = "https://edta-api.onrender.com/scenario-demo"
GITHUB = "https://github.com/jeraldcs/edta-complete-api"
SWAGGER = "https://edta-api.onrender.com/docs"
OPENAPI = "https://edta-api.onrender.com/openapi.json"

# Reuse styling helpers from architecture docx generator
sys.path.insert(0, str(ROOT))
from generate_architecture_docx import (  # noqa: E402
    add_bullet,
    add_code_block,
    add_hyperlink,
    add_image,
    add_table,
    configure_styles,
    draw_four_tier_diagram,
    draw_layer_stack_diagram,
    pdf_pages_to_png,
)


def ensure_diagram_pdf() -> Path:
    if FULL_PDF.exists():
        return FULL_PDF
    script = ROOT / "generate_architecture_diagrams_pdf.py"
    if script.exists():
        subprocess.run([sys.executable, str(script)], check=True, cwd=ROOT.parent)
    if FULL_PDF.exists():
        return FULL_PDF
    if not LEGACY_PDF.exists():
        import generate_architecture_pdf as gap

        gap.main()
    return LEGACY_PDF


def pdf_cover_to_png(pdf_path: Path, page_index: int, out_path: Path, dpi: int = 150) -> Path:
    out_path.parent.mkdir(parents=True, exist_ok=True)
    pdf = fitz.open(pdf_path)
    page = pdf[page_index]
    page.get_pixmap(dpi=dpi).save(str(out_path))
    pdf.close()
    return out_path


def add_metadata_links(doc: Document) -> None:
    for label, url in [
        ("Live demo", LIVE_DEMO),
        ("GitHub repository", GITHUB),
        ("OpenAPI / Swagger", SWAGGER),
    ]:
        p = doc.add_paragraph(style="Metadata")
        p.add_run(f"{label}: ")
        add_hyperlink(p, url, url)


def add_cover(doc: Document) -> None:
    doc.add_paragraph("EDTA Complete Showcase", style="Title")
    doc.add_paragraph(
        "Experience-Driven Targeting Architecture — Governed, Explainable, Trust-Aware Personalization",
        style="Doc Subtitle",
    )
    doc.add_paragraph("Author: Jerald Selvaraj · jerald.cs@gmail.com", style="Metadata")
    doc.add_paragraph("Version: Showcase v1 — July 2026 · Reference implementation v2.4", style="Metadata")
    add_metadata_links(doc)
    doc.add_paragraph()
    doc.add_paragraph(
        "This document is a single showcase package for architects, engineers, product leaders, "
        "and publication reviewers. It explains what EDTA is, how the live demo works, how to "
        "integrate the API into any application, why it benefits every industry, and where the "
        "project is headed.",
        style="Normal",
    )


def add_toc(doc: Document) -> None:
    doc.add_page_break()
    doc.add_paragraph("Table of Contents", style="Heading 1")
    sections = [
        "1. Executive Summary",
        "2. Why EDTA Is Useful for Everyone",
        "3. Problem Statement",
        "4. Traditional vs EDTA — Comparison Framework",
        "5. Architecture Overview",
        "6. Core Components & Code Map",
        "7. How the Live Demo Works",
        "8. API Integration Guide",
        "9. Benefits — Engineering, Business, Compliance",
        "10. Multi-Industry Applicability",
        "11. Benchmarks & Evidence",
        "12. Security, Deployment & Operations",
        "13. Future Roadmap",
        "14. Publication & InfoQ Materials",
        "15. Resources & Quick Start",
    ]
    for item in sections:
        doc.add_paragraph(item, style="List Number")


def section_executive_summary(doc: Document) -> None:
    doc.add_page_break()
    doc.add_paragraph("1. Executive Summary", style="Heading 1")
    doc.add_paragraph(
        "Enterprise teams are adopting large language models for personalization faster than they are "
        "adopting governance for personalization. The result is often relevant but risky recommendations: "
        "opaque scores, uncontrolled token cost, and weak answers when product, compliance, or trust teams "
        "ask why a recommendation was shown."
    )
    doc.add_paragraph(
        "EDTA (Experience-Driven Targeting Architecture) is an open-source reference implementation "
        "(FastAPI, Python) that decomposes personalization into explicit, auditable modules instead of "
        "a single black-box model."
    )
    add_table(
        doc,
        ["Acronym", "Name", "Role"],
        [
            ["TKGE", "Temporal Knowledge Graph Engine", "Journey timeline, inferred intent"],
            ["EML", "Experience Memory Layer", "Trust, fatigue, preferences, outcomes"],
            ["HAOE", "Hybrid AI Orchestration Engine", "Rules → SLM → ML → optional LLM"],
            ["EDS", "Experience DNA Score", "Explainable relevance breakdown"],
            ["TAPL", "Trust-Aware Personalization Layer", "show / soften / delay / suppress in ranking"],
            ["OSE", "Outcome Simulation Engine", "Conversion, revenue, trust, compliance risk"],
        ],
    )
    doc.add_paragraph(
        "The LLM enriches and explains; it does not own consent, compliance, fatigue, or final trust governance. "
        "A live browser demo, OpenAPI contract, 167 automated tests, and reproducible benchmarks are included."
    )


def section_why_useful(doc: Document) -> None:
    doc.add_paragraph("2. Why EDTA Is Useful for Everyone", style="Heading 1")
    audiences = [
        (
            "Software architects",
            "Reference pattern for multi-tier AI routing, policy externalization, and API-native explainability.",
        ),
        (
            "Backend / ML engineers",
            "Runnable code for intent, journey, TAPL, outcome simulation, and ranker modules with clear boundaries.",
        ),
        (
            "Product managers",
            "Demo shows before/after trust and fatigue, TAPL actions, and outcome scores — not only candidate IDs.",
        ),
        (
            "Compliance & trust teams",
            "TAPL modifies ranking; audit trails in SQLite; rules fired and tier visible in every response.",
        ),
        (
            "Industry vertical teams",
            "Swap YAML rule packs, candidate catalogs, and empathy profiles — keep the same engine core.",
        ),
        (
            "Researchers & publishers",
            "Reproducible four-tier benchmarks, architecture diagrams, and honest scope limits for articles.",
        ),
    ]
    for role, value in audiences:
        add_bullet(doc, value, bold_prefix=f"{role}: ")


def section_problem(doc: Document) -> None:
    doc.add_paragraph("3. Problem Statement", style="Heading 1")
    add_code_block(doc, "context → model → ranked recommendation")
    doc.add_paragraph("This pattern fails enterprise requirements:")
    add_table(
        doc,
        ["Gap", "Consequence"],
        [
            ["Opaque ranking", "Cannot answer why in production incidents"],
            ["LLM-first design", "Unbounded cost, latency, provider dependency"],
            ["Trust as metadata only", "Consent/fatigue logged but not enforced in ranking"],
            ["Single-channel thinking", "Web, chatbot, IoT need unified policy"],
            ["No outcome view", "Optimizes clicks, not expected business outcome"],
            ["Vertical monolith", "Each industry rewrites the engine instead of policy layers"],
        ],
    )


def section_traditional_vs_edta(doc: Document) -> None:
    doc.add_page_break()
    doc.add_paragraph("4. Traditional vs EDTA — Comparison Framework", style="Heading 1")
    doc.add_paragraph(
        "EDTA is a demo and reference implementation built from experience with traditional "
        "personalization architectures — rule engines, single-model rankers, and batch CRM campaigns — "
        "then extended with modern AI under explicit governance."
    )
    doc.add_paragraph(
        "EDTA does not replace traditional personalization — it completes it. Rules stay fast and "
        "auditable; ML stays for catalog fit; AI enters only where ambiguity demands it; and governance, "
        "memory, and explainability sit inside the decision path, not in a downstream dashboard."
    )

    doc.add_paragraph("4.1 Three traditional patterns", style="Heading 2")
    add_table(
        doc,
        ["Pattern", "Strengths", "Breaks when…"],
        [
            ["Rule-based personalization", "Auditable, fast, cheap", "Multi-channel NLP and journey memory are required"],
            ["Monolithic ranker", "Strong catalog match", "Consent/fatigue must change the rank, not only logs"],
            ["LLM bolt-on", "Handles ambiguity", "Cost, latency, and audit become first-class requirements"],
        ],
    )

    doc.add_paragraph("4.2 What you can prove in the live demo", style="Heading 2")
    add_table(
        doc,
        ["Dimension", "Traditional (typical)", "EDTA"],
        [
            ["Explainability", "Score only or post-hoc SHAP", "Rules, tier, TAPL, EDS, explanation in JSON"],
            ["Governance in rank", "Consent outside the ranker", "TAPL show/soften/delay/suppress/fallback"],
            ["Memory", "Stateless or batch segments", "EML trust/fatigue across sessions"],
            ["Outcome view", "Click-optimized rank", "OSE conversion/revenue/trust risk"],
            ["Cost control", "Fixed infra or unbounded LLM", "HAOE Rules → SLM → ML → optional LLM"],
            ["Vertical reuse", "Rewrite engine per industry", "Same core; swap YAML + catalog"],
        ],
    )

    doc.add_paragraph("4.3 Five-minute comparison script", style="Heading 2")
    for step in [
        "Rules-era upsell — structured SUV booking; show TAPL + EDS + explanation (chip: Rules-era upsell).",
        "ML ranker era — family vacation with cargo; Semantic alone is not final (chip: ML ranker era).",
        "LLM bolt-on risk — ambiguous one-liner; inference tier visible, TAPL still governs (chip: LLM bolt-on risk).",
        "EDTA governed — loyalty family SUV with consent; full stack panels (chip: EDTA governed).",
        "Governance contrast — No consent → generic_fallback; High fatigue → delay/soften.",
    ]:
        add_bullet(doc, step)

    doc.add_paragraph("4.4 Evidence and honesty bounds", style="Heading 2")
    for item in [
        "Travel matrix: docs/TRAVEL_SCENARIO_SCORES.md — alignment proxy, not production A/B lift.",
        "Four-tier benchmarks: docs/BENCHMARKS.md — use local Docker for engine latency; label Render cold start separately.",
        "Counterfactuals: POST /v1/compare-outcomes — outcome simulation before live experiments.",
        "Full exhibit: docs/TRADITIONAL_VS_EDTA.md — reusable for InfoQ and customer reviews.",
    ]:
        add_bullet(doc, item)


def section_architecture(doc: Document, layer_png: Path, tier_png: Path, diagram_pngs: list[Path]) -> None:
    doc.add_page_break()
    doc.add_paragraph("5. Architecture Overview", style="Heading 1")
    add_code_block(
        doc,
        """Customer Context
  → Profile Lookup + Experience Memory (EML)
  → Temporal Knowledge Graph (TKGE)
  → Hybrid AI Orchestration (HAOE) — Rules / SLM / ML / LLM
  → Intent + Journey Inference
  → Per-candidate: EDS → TAPL → OSE → Final Ranker
  → Explanation + Memory Update + Feedback Learning
  → Recommendation""",
    )
    add_image(doc, layer_png, 6.5, "Figure 1 — EDTA logical layer stack")
    add_image(doc, tier_png, 6.5, "Figure 2 — HAOE four-tier inference routing")
    captions = [
        "Figure 3 — System context (clients, API, external services)",
        "Figure 4 — Application container and code modules",
        "Figure 5 — End-to-end recommend sequence",
        "Figure 6 — Scenario demo path",
    ]
    for i, path in enumerate(diagram_pngs[:4]):
        cap = captions[i] if i < len(captions) else f"Architecture diagram {i + 3}"
        add_image(doc, path, 6.5, cap)


def section_code_map(doc: Document) -> None:
    doc.add_page_break()
    doc.add_paragraph("6. Core Components & Code Map", style="Heading 1")
    doc.add_paragraph("Repository layout (github.com/jeraldcs/edta-complete-api):")
    add_code_block(
        doc,
        """edta-complete-api/
├── app/
│   ├── main.py                 # FastAPI entry, middleware, routes
│   ├── api/v1/router.py        # Canonical /v1 API
│   ├── demo_proxy.py           # Server-side auth for browser demo
│   ├── container.py            # ServiceContainer wiring
│   ├── services/recommendation_handlers.py  # Request orchestration
│   ├── recommender.py          # RecommendationEngine rank loop
│   ├── orchestration.py        # HAOE tier selection
│   ├── experience_memory.py    # EML
│   ├── context_graph.py        # TKGE
│   ├── empathy/service.py      # Travel empathy engine
│   ├── inference/              # ML service, explanation router
│   ├── ai/                     # sklearn model wrappers
│   ├── llm/                    # Optional OpenAI / SLM clients
│   └── db.py                   # SQLite schema
├── config/                     # TAPL, HAOE, rules, empathy YAML
├── data/                       # Catalogs, training CSV, profiles
├── models/                     # Trained *.joblib artifacts
├── static/                     # scenario-demo UI
├── tests/                      # 167 pytest tests
└── scripts/                    # train, benchmark, load test""",
    )
    add_table(
        doc,
        ["Module", "File", "Responsibility"],
        [
            ["API handlers", "recommendation_handlers.py", "Enrich → rank → persist → respond"],
            ["Rank engine", "recommender.py", "Candidate loop, tier execution, hybrid score"],
            ["HAOE", "orchestration.py", "Rules/SLM/ML/LLM routing, cost budget"],
            ["TAPL", "ai/tapl_model.py + tapl_policies.yaml", "Trust governance actions"],
            ["EML", "experience_memory.py", "Trust, fatigue, preferences, history"],
            ["TKGE", "context_graph.py + graph_store.py", "Temporal journey graph"],
            ["Empathy", "empathy/service.py", "Hidden needs, TCO, constraints (travel)"],
        ],
    )


def section_demo(doc: Document) -> None:
    doc.add_page_break()
    doc.add_paragraph("7. How the Live Demo Works", style="Heading 1")
    p = doc.add_paragraph()
    p.add_run("URL: ")
    add_hyperlink(p, LIVE_DEMO, LIVE_DEMO)

    doc.add_paragraph("7.1 User flow", style="Heading 2")
    steps = [
        "Open /scenario-demo in a browser (Render may cold-start ~30s on free tier).",
        "Select a trained scenario chip or type free-text (e.g. family SUV airport rental).",
        "Optional: use Compare vs traditional chips (Rules-era, ML ranker, LLM bolt-on, EDTA governed) — same Recommend API.",
        "Click Recommend — demo-api.js calls the API (production uses /demo-api proxy with server-side auth).",
        "View top recommendation with EDS, TAPL, outcome scores, and explanation.",
        "Architecture panels show TKGE, EML, HAOE, and OSE summaries from the response JSON.",
        "Optional: scroll the 11-scenario travel benchmark matrix; click a row to load and run that scenario.",
        "Feedback buttons (Click / Convert / Dismiss) update EML and TKGE for the scenario subject.",
    ]
    for step in steps:
        add_bullet(doc, step)

    doc.add_paragraph("7.2 Server-side demo path", style="Heading 2")
    add_code_block(
        doc,
        """scenario_text
  → ScenarioNLPParser (channel, intent, journey, consent hints)
  → scenario_anonymous_id() — isolates EML/TKGE per trained profile
  → ProfileLookupService (customer_id → CRM mock JSON)
  → ExperienceMemoryLayer enrich + calibration
  → EmpathyEngine.process() — hidden needs, TCO, constraints
  → Merge vehicle catalog candidates
  → RecommendationEngine.recommend()
  → JSON response + architecture panels""",
    )

    doc.add_paragraph("7.3 Demo modes", style="Heading 2")
    add_table(
        doc,
        ["Mode", "Endpoint", "Purpose"],
        [
            ["Recommend", "POST /recommend-from-scenario", "Full stack ranking from free text"],
            ["Simulate", "POST /simulate", "Outcome simulation without full rank response"],
            ["Experience memory", "GET /experience-memory", "Trust/fatigue snapshot for subject"],
            ["Travel benchmark", "GET /travel-scenario-benchmark", "11-scenario score matrix"],
            ["Empathy", "POST /empathy/simulate", "Hidden needs + TCO bundle"],
        ],
    )


def section_api_integration(doc: Document) -> None:
    doc.add_page_break()
    doc.add_paragraph("8. API Integration Guide", style="Heading 1")
    doc.add_paragraph(
        "New integrations should use the versioned /v1 API. Legacy root routes remain for the web demo."
    )

    doc.add_paragraph("8.1 Authentication", style="Heading 2")
    doc.add_paragraph(
        "When EDTA_API_KEY is set, send X-API-Key on mutating and protected read requests. "
        "Browser demos on production use /demo-api/* proxy (same origin, no client key)."
    )

    doc.add_paragraph("8.2 Core endpoints", style="Heading 2")
    add_table(
        doc,
        ["Method", "Path", "Purpose"],
        [
            ["POST", "/v1/recommend", "Structured CustomerContext + optional candidates"],
            ["POST", "/v1/recommend-from-scenario", "Free-text scenario (demo-friendly)"],
            ["POST", "/v1/recommend/batch", "Async batch (202 + job polling)"],
            ["POST", "/v1/feedback", "Click / convert / dismiss → EML update"],
            ["POST", "/v1/compare-outcomes", "Counterfactual A vs B"],
            ["GET", "/v1/context-graph", "TKGE export for subject"],
            ["GET", "/v1/experience-memory", "EML snapshot"],
            ["GET", "/v1/orchestration-status", "HAOE telemetry"],
            ["POST", "/v1/webhooks", "Register HTTPS webhook (SSRF-validated)"],
        ],
    )

    doc.add_paragraph("8.3 Example — structured recommend (curl)", style="Heading 2")
    add_code_block(
        doc,
        """curl -X POST "https://your-host/v1/recommend" \\
  -H "Content-Type: application/json" \\
  -H "X-API-Key: YOUR_EDTA_API_KEY" \\
  -d @sample_requests/web_family_suv.json""",
    )

    doc.add_paragraph("8.4 Example — scenario recommend (JavaScript fetch)", style="Heading 2")
    add_code_block(
        doc,
        """const response = await fetch("https://your-host/v1/recommend-from-scenario", {
  method: "POST",
  headers: {
    "Content-Type": "application/json",
    "X-API-Key": process.env.EDTA_API_KEY,
  },
  body: JSON.stringify({
    scenario_text: "Family of four needs SUV for Orlando theme park trip with luggage.",
    limit: 3,
    use_ai_models: true,
    include_empathy: true,
  }),
});
const data = await response.json();
// data.request_summary.inference.tier — rules | slm | ml | llm
// data.recommendations[0].ai_score.tapl.action — show | soften | delay | suppress
// data.recommendations[0].explanation — human-readable rationale""",
    )

    doc.add_paragraph("8.5 Integration patterns", style="Heading 2")
    patterns = [
        ("Web / mobile app", "Call POST /v1/recommend from BFF with enriched CustomerContext from your CDP."),
        ("Chatbot / voice", "Use POST /recommend-from-scenario for NLP input; set channel to chatbot in parsed context."),
        ("Partner API", "Use /v1 with API key; poll /v1/jobs/{id} for batch campaigns."),
        ("Event-driven", "Register POST /v1/webhooks; subscribe to recommendation.created and feedback.received."),
        ("Idempotency", "Send Idempotency-Key header on POST /v1/recommend and /v1/feedback for safe retries."),
    ]
    for name, desc in patterns:
        add_bullet(doc, desc, bold_prefix=f"{name}: ")

    doc.add_paragraph("8.6 Response fields integrators should persist", style="Heading 2")
    add_bullet(doc, "request_summary.request_id — correlate logs and support tickets")
    add_bullet(doc, "request_summary.inference.tier — which AI tier ran")
    add_bullet(doc, "recommendations[].ai_score.tapl.action — governance decision")
    add_bullet(doc, "recommendations[].eds_score — explainable relevance breakdown")
    add_bullet(doc, "recommendations[].explanation_source — local | slm | llm")


def section_benefits(doc: Document) -> None:
    doc.add_page_break()
    doc.add_paragraph("9. Benefits — Engineering, Business, Compliance", style="Heading 1")
    doc.add_paragraph("9.1 Engineering benefits", style="Heading 2")
    for item in [
        "Modular testability — 167 pytest tests; force inference_mode per tier for benchmarks.",
        "Cost control — HAOE session budget, LLM circuit breaker, SLM distilled memory.",
        "Provider independence — runs fully without OpenAI; LLM is optional enrichment.",
        "API-native explainability — no log archaeology for why a recommendation was shown.",
        "Policy as code — TAPL, HAOE, rules in YAML without redeploying ranker code.",
    ]:
        add_bullet(doc, item)

    doc.add_paragraph("9.2 Business benefits", style="Heading 2")
    for item in [
        "Outcome-oriented ranking — OSE estimates conversion and revenue, not only similarity.",
        "Fatigue-aware engagement — EML + TAPL reduce over-exposure and banner blindness.",
        "Cross-channel consistency — same engine for web, mobile, chatbot, IoT, partner API.",
        "Faster vertical experiments — swap catalog + YAML; keep orchestration and memory layers.",
        "Empathy-aware travel upsell — TCO and hidden-needs matching for rental scenarios.",
    ]:
        add_bullet(doc, item)

    doc.add_paragraph("9.3 Compliance & trust benefits", style="Heading 2")
    for item in [
        "Consent enforcement — TAPL generic_fallback when personalization consent is false.",
        "Sensitive channel governance — suppress/soften on SMS, wearable, IoT for high-compliance content.",
        "Audit trail — TAPL decisions logged to SQLite with reason strings.",
        "Healthcare education scope — HCP content routing with disclaimers; not clinical DSS.",
        "Problem+JSON errors with request_id — production-safe error correlation.",
    ]:
        add_bullet(doc, item)


def section_industries(doc: Document) -> None:
    doc.add_page_break()
    doc.add_paragraph("10. Multi-Industry Applicability", style="Heading 1")
    doc.add_paragraph(
        "EDTA is vertical-agnostic at the engine layer and vertical-specific at the policy and catalog layer. "
        "The same four-tier pipeline applies; teams supply domain rules, approved candidates, and empathy profiles."
    )
    add_table(
        doc,
        ["Industry", "Example candidates", "EDTA mechanisms", "Demo scenario"],
        [
            ["Travel / car rental", "SUV upgrade, booking discount", "Empathy, TCO, TKGE", "11 road-trip scenarios"],
            ["Hotel / hospitality", "Room offer, reservation assist", "Rules + journey stage", "hotel_reservation_assist"],
            ["Restaurant", "QR menu, dish promo", "Channel fit, mobile", "mobile_restaurant_qr_menu"],
            ["Banking", "Product offer, advisor content", "TAPL compliance sensitivity", "banking rules pack"],
            ["Healthcare", "HCP education content", "TAPL suppress on SMS/wearable", "obesity_product_hcp_education"],
            ["Retail / e-commerce", "Upsell, retention offer", "EML preferences + OSE", "catalog JSON swap"],
            ["IoT / connected car", "Location assist, service nudge", "Channel=connected_car", "connected_car_location_assist"],
        ],
    )
    doc.add_paragraph("How to adapt for a new industry (checklist):", style="Heading 2")
    for step in [
        "Add candidates to data/catalog/ or pass candidates in the API request.",
        "Create or extend config/rules/packs/{industry}.yaml.",
        "Review config/tapl_policies.yaml for consent, fatigue, and channel rules.",
        "Add training rows to data/scenario_training_master.csv; run scripts/train_all_models.py.",
        "Optional: add empathy or enrichment profiles under config/empathy/.",
        "Point PROFILE_LOOKUP_ADAPTER at your CDP/CRM instead of json_file.",
    ]:
        add_bullet(doc, step)


def section_benchmarks(doc: Document) -> None:
    doc.add_page_break()
    doc.add_paragraph("11. Benchmarks & Evidence", style="Heading 1")
    doc.add_paragraph("11.1 Travel scenario matrix (11 scenarios)", style="Heading 2")
    doc.add_paragraph(
        "100% top-rank vehicle alignment on the demo training set (alignment proxy — not production A/B lift). "
        "See docs/TRAVEL_SCENARIO_SCORES.md and GET /travel-scenario-benchmark."
    )
    doc.add_paragraph("11.2 Four-tier benchmarks", style="Heading 2")
    add_table(
        doc,
        ["Tier", "p95 target (local)", "Cost", "Use case"],
        [
            ["Rules", "< 50 ms", "~0 units", "Structured context, auditable"],
            ["SLM", "< 100 ms", "Low", "Distilled pattern memory"],
            ["ML", "< 300 ms", "Medium", "sklearn classifiers"],
            ["LLM", "3–8 s", "Highest", "Ambiguous rich context (optional)"],
        ],
    )
    add_code_block(doc, "python scripts/load_test_tiers.py --base-url http://127.0.0.1:8000")
    doc.add_paragraph("11.3 Test suite", style="Heading 2")
    add_code_block(doc, "pytest   # 167 tests — API, TAPL, empathy, HAOE, security, scenarios")


def section_ops(doc: Document) -> None:
    doc.add_page_break()
    doc.add_paragraph("12. Security, Deployment & Operations", style="Heading 1")
    add_table(
        doc,
        ["Topic", "Implementation"],
        [
            ["Authentication", "X-API-Key when EDTA_API_KEY set; demo proxy in production"],
            ["CORS", "Explicit allowlist in production"],
            ["Webhooks", "HTTPS only; SSRF validation on registration"],
            ["Rate limiting", "Per-IP in memory; use edge limiter at scale"],
            ["Health probes", "/live (liveness), /ready (readiness), /metrics (Prometheus)"],
            ["Storage", "SQLite (EML, TKGE, audit); mount persistent disk for production"],
            ["Deploy", "Docker, Render blueprint (render.yaml), docker-compose locally"],
        ],
    )
    doc.add_paragraph("Production checklist:", style="Heading 2")
    for item in [
        "Set EDTA_API_KEY and EDTA_ENVIRONMENT=production",
        "Set EDTA_CORS_ORIGINS to your app origin",
        "Use /ready for load balancer health checks",
        "Enable EDTA_LOG_FORMAT=json and scrape /metrics",
        "Review config/tapl_policies.yaml for your regulatory context",
    ]:
        add_bullet(doc, item)


def section_roadmap(doc: Document) -> None:
    doc.add_page_break()
    doc.add_paragraph("13. Future Roadmap", style="Heading 1")
    doc.add_paragraph("Completed (v2.4 reference stack)", style="Heading 2")
    for item in [
        "Four-tier HAOE with benchmarks and distilled SLM memory",
        "EML + TKGE with feedback loop",
        "Empathy engine — hidden needs, TCO, 11 travel profiles (Phase 6A–6B largely complete)",
        "Production hardening — auth on reads, demo proxy, webhook SSRF, idempotency, SQLite feedback",
        "167-test CI with ML eval and OpenAPI export",
    ]:
        add_bullet(doc, item)

    doc.add_paragraph("Near-term (Phase 6C–7)", style="Heading 2")
    add_table(
        doc,
        ["Phase", "Focus", "Outcome"],
        [
            ["6C", "Live weather/route/gas enrichment APIs", "AWD boost with forecast evidence"],
            ["6D", "Dedicated empathy simulator UI polish", "Enhanced public demo UX"],
            ["7", "Empathy accuracy evaluation vs baseline filters", "Published benchmark exhibit"],
            ["Ops", "PostgreSQL adapter option", "Multi-instance production deploy"],
            ["Ops", "Redis rate limiting + webhook retry DLQ", "Scale-out readiness"],
        ],
    )

    doc.add_paragraph("Medium-term", style="Heading 2")
    for item in [
        "HTTP CDP/CRM profile adapter (replace json_file mock)",
        "Partner production case study with A/B metrics",
        "Kubernetes Helm chart and horizontal pod autoscaling guide",
        "Optional federated learning hook for ranker retraining",
        "Expanded healthcare and banking rule packs with compliance review templates",
    ]:
        add_bullet(doc, item)


def section_publication(doc: Document) -> None:
    doc.add_page_break()
    doc.add_paragraph("14. Publication & InfoQ Materials", style="Heading 1")
    doc.add_paragraph(
        "Supporting documents in the docs/ folder for architecture publication and InfoQ submission:"
    )
    pubs = [
        ("EDTA_PUBLICATION_DOCUMENT.md", "Full publication package"),
        ("EDTA_ARCHITECTURE_DOCUMENT.md", "Detailed architecture reference"),
        ("edta_architecture_diagrams_full.pdf", "11 architecture diagrams"),
        ("TRADITIONAL_VS_EDTA.md", "Traditional vs EDTA comparison exhibit"),
        ("INFOQ_EDITOR_PITCH.md", "Editor submission email template"),
        ("INFOQ_ARTICLE_OUTLINE.md", "~3,000-word article outline"),
        ("edta_llm_personalization_article_draft.md", "Extended article draft"),
        ("EDTA_POSITIONING.md", "Travel vs healthcare narrative angles"),
    ]
    for name, desc in pubs:
        add_bullet(doc, desc, bold_prefix=f"{name}: ")

    doc.add_paragraph("Suggested working title for publication:", style="Heading 2")
    doc.add_paragraph(
        "Governed Personalization: A Four-Tier Architecture for Explainable, Trust-Aware Recommendations",
        style="Doc Subtitle",
    )


def section_resources(doc: Document) -> None:
    doc.add_page_break()
    doc.add_paragraph("15. Resources & Quick Start", style="Heading 1")
    add_metadata_links(doc)
    doc.add_paragraph()
    doc.add_paragraph("Local quick start:", style="Heading 2")
    add_code_block(
        doc,
        """git clone https://github.com/jeraldcs/edta-complete-api.git
cd edta-complete-api
python -m venv .venv
.venv\\Scripts\\activate          # Windows
pip install -r requirements.txt
python scripts/train_all_models.py
uvicorn app.main:app --reload --port 8000
# Demo: http://localhost:8000/scenario-demo
pytest""",
    )
    doc.add_paragraph("Regenerate this showcase document:", style="Heading 2")
    add_code_block(doc, "python docs/generate_showcase_docx.py")
    doc.add_paragraph("Regenerate architecture diagram PDF:", style="Heading 2")
    add_code_block(doc, "python docs/generate_architecture_diagrams_pdf.py")


def build() -> None:
    ASSETS.mkdir(parents=True, exist_ok=True)
    pdf_path = ensure_diagram_pdf()
    layer_png = draw_layer_stack_diagram(ASSETS / "layer_stack.png")
    tier_png = draw_four_tier_diagram(ASSETS / "four_tier.png")

    diagram_pngs: list[Path] = []
    if pdf_path == FULL_PDF:
        for i in range(1, 5):
            out = ASSETS / f"full_diagram_{i}.png"
            pdf_cover_to_png(FULL_PDF, i, out)
            diagram_pngs.append(out)
    else:
        diagram_pngs = pdf_pages_to_png(pdf_path, ASSETS)[:3]

    doc = Document()
    configure_styles(doc)
    add_cover(doc)
    add_toc(doc)
    section_executive_summary(doc)
    section_why_useful(doc)
    section_problem(doc)
    section_traditional_vs_edta(doc)
    section_architecture(doc, layer_png, tier_png, diagram_pngs)
    section_code_map(doc)
    section_demo(doc)
    section_api_integration(doc)
    section_benefits(doc)
    section_industries(doc)
    section_benchmarks(doc)
    section_ops(doc)
    section_roadmap(doc)
    section_publication(doc)
    section_resources(doc)

    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run(f"EDTA Complete Showcase — {GITHUB}")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    doc.save(OUTPUT)
    print(OUTPUT.resolve())


if __name__ == "__main__":
    build()
