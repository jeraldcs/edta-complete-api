from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "edta_llm_personalization_article_draft.md"
OUTPUT = ROOT / "edta_llm_personalization_article_draft.docx"


BLUE = RGBColor(0x2E, 0x74, 0xB5)
DARK_BLUE = RGBColor(0x1F, 0x4D, 0x78)
MUTED = RGBColor(0x55, 0x55, 0x55)
INK = RGBColor(0x00, 0x00, 0x00)
CODE_FILL = "F4F6F9"


def set_spacing(style, before=0, after=8, line=1.333):
    fmt = style.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def configure_styles(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    set_spacing(normal, before=0, after=8, line=1.333)

    title = doc.styles["Title"]
    title.font.name = "Calibri"
    title.font.size = Pt(22)
    title.font.bold = True
    title.font.color.rgb = DARK_BLUE
    set_spacing(title, before=0, after=8, line=1.15)

    subtitle = doc.styles.add_style("Article Subtitle", WD_STYLE_TYPE.PARAGRAPH)
    subtitle.font.name = "Calibri"
    subtitle.font.size = Pt(11)
    subtitle.font.italic = True
    subtitle.font.color.rgb = MUTED
    set_spacing(subtitle, before=0, after=12, line=1.2)

    for style_name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        set_spacing(style, before=before, after=after, line=1.2)

    bullet = doc.styles["List Bullet"]
    bullet.font.name = "Calibri"
    bullet.font.size = Pt(11)
    set_spacing(bullet, before=0, after=4, line=1.208)
    bullet.paragraph_format.left_indent = Inches(0.375)
    bullet.paragraph_format.first_line_indent = Inches(-0.194)

    number = doc.styles["List Number"]
    number.font.name = "Calibri"
    number.font.size = Pt(11)
    set_spacing(number, before=0, after=4, line=1.208)
    number.paragraph_format.left_indent = Inches(0.375)
    number.paragraph_format.first_line_indent = Inches(-0.194)

    code = doc.styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    code.font.name = "Courier New"
    code.font.size = Pt(8.5)
    code.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
    set_spacing(code, before=4, after=6, line=1.0)
    code.paragraph_format.left_indent = Inches(0.18)
    code.paragraph_format.right_indent = Inches(0.18)

    meta = doc.styles.add_style("Metadata", WD_STYLE_TYPE.PARAGRAPH)
    meta.font.name = "Calibri"
    meta.font.size = Pt(10)
    meta.font.color.rgb = MUTED
    set_spacing(meta, before=0, after=3, line=1.15)


def shade_paragraph(paragraph, fill: str):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


def add_horizontal_rule(doc: Document):
    p = doc.add_paragraph()
    p_format = p.paragraph_format
    p_format.space_before = Pt(3)
    p_format.space_after = Pt(9)
    p_pr = p._p.get_or_add_pPr()
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "DADCE0")
    border.append(bottom)
    p_pr.append(border)


def add_runs_from_inline_markdown(paragraph, text: str):
    pattern = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            paragraph.add_run(text[pos:match.start()])
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Courier New"
            run.font.size = Pt(9.5)
        pos = match.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def emit_paragraph(doc: Document, text: str):
    if not text:
        return
    p = doc.add_paragraph()
    add_runs_from_inline_markdown(p, text)


def convert():
    doc = Document()
    configure_styles(doc)

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    in_code = False
    code_lines = []
    paragraph_buffer = []

    def flush_paragraph():
        nonlocal paragraph_buffer
        if paragraph_buffer:
            emit_paragraph(doc, " ".join(paragraph_buffer).strip())
            paragraph_buffer = []

    for raw in lines:
        line = raw.rstrip()

        if line.startswith("```"):
            if in_code:
                for code_line in code_lines:
                    p = doc.add_paragraph(code_line if code_line else " ", style="Code Block")
                    shade_paragraph(p, CODE_FILL)
                code_lines = []
                in_code = False
            else:
                flush_paragraph()
                in_code = True
            continue

        if in_code:
            code_lines.append(line)
            continue

        if not line.strip():
            flush_paragraph()
            continue

        if line == "---":
            flush_paragraph()
            add_horizontal_rule(doc)
            continue

        if line.startswith("# "):
            flush_paragraph()
            p = doc.add_paragraph(line[2:].strip(), style="Title")
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            continue

        if line.startswith("## "):
            flush_paragraph()
            text = line[3:].strip()
            style = "Heading 1" if not text.startswith("Draft For") else "Article Subtitle"
            doc.add_paragraph(text, style=style)
            continue

        if line.startswith("### "):
            flush_paragraph()
            doc.add_paragraph(line[4:].strip(), style="Heading 2")
            continue

        if line.startswith("- "):
            flush_paragraph()
            p = doc.add_paragraph(style="List Bullet")
            add_runs_from_inline_markdown(p, line[2:].strip())
            continue

        number_match = re.match(r"^(\d+)\.\s+(.*)$", line)
        if number_match:
            flush_paragraph()
            p = doc.add_paragraph(style="List Number")
            add_runs_from_inline_markdown(p, number_match.group(2).strip())
            continue

        if line.startswith("**") and line.endswith("**") and line.count("**") == 2:
            flush_paragraph()
            p = doc.add_paragraph(style="Metadata")
            add_runs_from_inline_markdown(p, line)
            continue

        paragraph_buffer.append(line)

    flush_paragraph()

    section = doc.sections[0]
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("EDTA personalization architecture article draft").font.size = Pt(8)

    doc.save(OUTPUT)
    print(OUTPUT.resolve())


if __name__ == "__main__":
    convert()
