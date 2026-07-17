"""
Generate EDTA publication Word document with embedded architecture diagrams.

Output: docs/EDTA_PUBLICATION_DOCUMENT.docx
Source:  docs/EDTA_PUBLICATION_DOCUMENT.md
Diagrams: docs/edta_architecture_diagram.pdf (3 pages, July 2026 codebase)
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

import fitz  # PyMuPDF
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "EDTA_PUBLICATION_DOCUMENT.md"
PDF_DIAGRAM = ROOT / "edta_architecture_diagram.pdf"
ASSETS = ROOT / "_publication_docx_assets"
OUTPUT = ROOT / "EDTA_PUBLICATION_DOCUMENT.docx"

LIVE_DEMO = "https://edta-api.onrender.com/scenario-demo"
GITHUB = "https://github.com/jeraldcs/edta-complete-api"
SWAGGER = "https://edta-api.onrender.com/docs"

BLUE = RGBColor(0x2E, 0x74, 0xB5)
DARK_BLUE = RGBColor(0x1F, 0x4D, 0x78)
MUTED = RGBColor(0x55, 0x55, 0x55)
INK = RGBColor(0x00, 0x00, 0x00)
CODE_FILL = "F4F6F9"


def set_spacing(style, before=0, after=8, line=1.333):
    fmt = style.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def configure_styles(doc: Document):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    set_spacing(normal)

    title = doc.styles["Title"]
    title.font.name = "Calibri"
    title.font.size = Pt(22)
    title.font.bold = True
    title.font.color.rgb = DARK_BLUE
    set_spacing(title, after=6)

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        st = doc.styles[name]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = color
        set_spacing(st, before=before, after=after)

    if "Doc Subtitle" not in [s.name for s in doc.styles]:
        sub = doc.styles.add_style("Doc Subtitle", WD_STYLE_TYPE.PARAGRAPH)
    else:
        sub = doc.styles["Doc Subtitle"]
    sub.font.name = "Calibri"
    sub.font.size = Pt(12)
    sub.font.italic = True
    sub.font.color.rgb = MUTED

    code = doc.styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH) if "Code Block" not in [
        s.name for s in doc.styles
    ] else doc.styles["Code Block"]
    code.font.name = "Consolas"
    code.font.size = Pt(8.5)
    set_spacing(code, before=4, after=6, line=1.0)

    meta = doc.styles.add_style("Metadata", WD_STYLE_TYPE.PARAGRAPH) if "Metadata" not in [
        s.name for s in doc.styles
    ] else doc.styles["Metadata"]
    meta.font.name = "Calibri"
    meta.font.size = Pt(10)
    meta.font.color.rgb = MUTED

    try:
        caption = doc.styles.add_style("Caption", WD_STYLE_TYPE.PARAGRAPH)
    except ValueError:
        caption = doc.styles["Caption"]
    caption.font.name = "Calibri"
    caption.font.size = Pt(9)
    caption.font.italic = True
    caption.font.color.rgb = MUTED
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER


def shade_paragraph(paragraph, fill: str):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


def add_hyperlink(paragraph, text: str, url: str):
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(color)
    r_pr.append(underline)
    run.append(r_pr)
    text_elem = OxmlElement("w:t")
    text_elem.text = text
    run.append(text_elem)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_runs_from_inline_markdown(paragraph, text: str):
    pattern = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            paragraph.add_run(text[pos : match.start()])
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        pos = match.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def emit_paragraph(doc: Document, text: str):
    if not text.strip():
        return
    p = doc.add_paragraph()
    add_runs_from_inline_markdown(p, text.strip())


def add_horizontal_rule(doc: Document):
    p = doc.add_paragraph()
    p_pr = p._p.get_or_add_pPr()
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "DADCE0")
    border.append(bottom)
    p_pr.append(border)


def add_image(doc: Document, path: Path, width_in: float = 6.8, caption: str | None = None):
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width_in))
    if caption:
        doc.add_paragraph(caption, style="Caption")


def add_table_from_markdown(doc: Document, header_line: str, row_lines: list[str]):
    headers = [c.strip() for c in header_line.strip("|").split("|")]
    rows = []
    for line in row_lines:
        if not line.strip().startswith("|"):
            continue
        cells = [c.strip() for c in line.strip("|").split("|")]
        if len(cells) == len(headers):
            rows.append(cells)
    if not headers or not rows:
        return
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
        for p in table.rows[0].cells[i].paragraphs:
            for r in p.runs:
                r.bold = True
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            table.rows[r_idx + 1].cells[c_idx].text = val
    doc.add_paragraph()


def pdf_pages_to_png(pdf_path: Path, out_dir: Path, dpi: int = 160) -> list[Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    paths: list[Path] = []
    pdf = fitz.open(pdf_path)
    for i, page in enumerate(pdf):
        pix = page.get_pixmap(dpi=dpi)
        path = out_dir / f"publication_diagram_page_{i + 1}.png"
        pix.save(str(path))
        paths.append(path)
    pdf.close()
    return paths


def ensure_pdf():
    if not PDF_DIAGRAM.exists():
        sys.path.insert(0, str(ROOT))
        import generate_architecture_pdf as gap

        gap.main()


def convert_markdown(doc: Document, lines: list[str], diagram_paths: list[Path]):
    in_code = False
    code_lines: list[str] = []
    paragraph_buffer: list[str] = []
    table_header: str | None = None
    table_rows: list[str] = []
    in_table = False
    diagram_captions = [
        "Figure 1 — Scenario demo runtime (scenario-demo → recommend-from-scenario → empathy → ranker → panels)",
        "Figure 2 — Codebase layer map (app/, config/, static/, empathy and AI modules)",
        "Figure 3 — HAOE four-tier inference + empathy pipeline",
    ]
    diagrams_inserted = False

    def flush_paragraph():
        nonlocal paragraph_buffer
        if paragraph_buffer:
            emit_paragraph(doc, " ".join(paragraph_buffer))
            paragraph_buffer = []

    def flush_table():
        nonlocal table_header, table_rows, in_table
        if table_header and table_rows:
            add_table_from_markdown(doc, table_header, table_rows)
        table_header = None
        table_rows = []
        in_table = False

    def insert_diagrams():
        nonlocal diagrams_inserted
        if diagrams_inserted:
            return
        doc.add_paragraph("Embedded architecture diagrams (July 2026 codebase)", style="Heading 3")
        for i, path in enumerate(diagram_paths):
            cap = diagram_captions[i] if i < len(diagram_captions) else f"Figure {i + 1}"
            add_image(doc, path, 6.75, cap)
        diagrams_inserted = True

    for raw in lines:
        line = raw.rstrip()

        if line.startswith("```"):
            flush_table()
            if in_code:
                for code_line in code_lines:
                    p = doc.add_paragraph(code_line if code_line else " ", style="Code Block")
                    shade_paragraph(p, CODE_FILL)
                code_lines = []
                in_code = False
            else:
                flush_paragraph()
                in_code = True
            continue

        if in_code:
            code_lines.append(line)
            continue

        if line.startswith("|") and "---" in line and table_header is None:
            continue

        if line.startswith("|"):
            flush_paragraph()
            if table_header is None:
                table_header = line
                in_table = True
            else:
                table_rows.append(line)
            continue
        elif in_table:
            flush_table()

        if not line.strip():
            flush_paragraph()
            continue

        if line == "---":
            flush_paragraph()
            add_horizontal_rule(doc)
            continue

        if line.startswith("# "):
            flush_paragraph()
            doc.add_paragraph(line[2:].strip(), style="Title")
            continue

        if line.startswith("## "):
            flush_paragraph()
            heading = line[3:].strip()
            doc.add_paragraph(heading, style="Heading 1")
            if heading.startswith("19. Architecture diagrams"):
                insert_diagrams()
            continue

        if line.startswith("### "):
            flush_paragraph()
            doc.add_paragraph(line[4:].strip(), style="Heading 2")
            continue

        if line.startswith("- "):
            flush_paragraph()
            p = doc.add_paragraph(style="List Bullet")
            add_runs_from_inline_markdown(p, line[2:].strip())
            continue

        number_match = re.match(r"^(\d+)\.\s+(.*)$", line)
        if number_match:
            flush_paragraph()
            p = doc.add_paragraph(style="List Number")
            add_runs_from_inline_markdown(p, number_match.group(2).strip())
            continue

        if "**" in line and line.strip().startswith("**"):
            flush_paragraph()
            p = doc.add_paragraph(style="Metadata")
            add_runs_from_inline_markdown(p, line)
            continue

        paragraph_buffer.append(line)

    flush_paragraph()
    flush_table()
    if not diagrams_inserted:
        insert_diagrams()


def build():
    ensure_pdf()
    diagram_paths = pdf_pages_to_png(PDF_DIAGRAM, ASSETS)

    doc = Document()
    configure_styles(doc)

    # Cover links (before markdown body)
    p = doc.add_paragraph(style="Metadata")
    p.add_run("Live demo: ")
    add_hyperlink(p, LIVE_DEMO, LIVE_DEMO)
    p = doc.add_paragraph(style="Metadata")
    p.add_run("GitHub: ")
    add_hyperlink(p, GITHUB, GITHUB)
    p = doc.add_paragraph(style="Metadata")
    p.add_run("OpenAPI: ")
    add_hyperlink(p, SWAGGER, SWAGGER)
    doc.add_paragraph()

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    convert_markdown(doc, lines, diagram_paths)

    section = doc.sections[0]
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run(f"EDTA Publication — {GITHUB}")
    run.font.size = Pt(8)
    run.font.color.rgb = MUTED

    doc.save(OUTPUT)
    print(OUTPUT.resolve())


if __name__ == "__main__":
    build()
