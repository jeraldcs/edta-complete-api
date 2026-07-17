"""
Generate EDTA Architecture Word document for GitHub / InfoQ publication.

Output: docs/EDTA_ARCHITECTURE_DOCUMENT.docx
Diagrams: embeds pages from edta_architecture_diagram.pdf + generated layer-stack PNG.
"""

from __future__ import annotations

from pathlib import Path

import fitz  # PyMuPDF
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parent
PDF_DIAGRAM = ROOT / "edta_architecture_diagram.pdf"
ASSETS = ROOT / "_docx_assets"
OUTPUT = ROOT / "EDTA_ARCHITECTURE_DOCUMENT.docx"

LIVE_DEMO = "https://edta-api.onrender.com/scenario-demo"
GITHUB = "https://github.com/jeraldcs/edta-complete-api"
SWAGGER = "https://edta-api.onrender.com/docs"

BLUE = RGBColor(0x2E, 0x74, 0xB5)
DARK_BLUE = RGBColor(0x1F, 0x4D, 0x78)
MUTED = RGBColor(0x55, 0x55, 0x55)
INK = RGBColor(0x00, 0x00, 0x00)
CODE_FILL = "F4F6F9"
LINK = RGBColor(0x05, 0x63, 0xC1)


def set_spacing(style, before=0, after=8, line=1.333):
    fmt = style.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def configure_styles(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    for margin in (section.top_margin, section.bottom_margin, section.left_margin, section.right_margin):
        pass
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    set_spacing(normal)

    title = doc.styles["Title"]
    title.font.name = "Calibri"
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = DARK_BLUE
    set_spacing(title, after=6)

    subtitle = doc.styles.add_style("Doc Subtitle", WD_STYLE_TYPE.PARAGRAPH)
    subtitle.font.name = "Calibri"
    subtitle.font.size = Pt(12)
    subtitle.font.italic = True
    subtitle.font.color.rgb = MUTED
    set_spacing(subtitle, after=10)

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        st = doc.styles[name]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = color
        set_spacing(st, before=before, after=after)

    bullet = doc.styles["List Bullet"]
    bullet.font.name = "Calibri"
    bullet.font.size = Pt(11)
    set_spacing(bullet, after=4)

    code = doc.styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    code.font.name = "Consolas"
    code.font.size = Pt(8.5)
    set_spacing(code, before=4, after=6, line=1.0)

    meta = doc.styles.add_style("Metadata", WD_STYLE_TYPE.PARAGRAPH)
    meta.font.name = "Calibri"
    meta.font.size = Pt(10)
    meta.font.color.rgb = MUTED
    set_spacing(meta, after=3)

    try:
        caption = doc.styles.add_style("Caption", WD_STYLE_TYPE.PARAGRAPH)
    except ValueError:
        caption = doc.styles["Caption"]
    caption.font.name = "Calibri"
    caption.font.size = Pt(9)
    caption.font.italic = True
    caption.font.color.rgb = MUTED
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_spacing(caption, before=2, after=10)


def shade_paragraph(paragraph, fill: str):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


def add_hyperlink(paragraph, text: str, url: str):
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(color)
    r_pr.append(underline)
    run.append(r_pr)
    text_elem = OxmlElement("w:t")
    text_elem.text = text
    run.append(text_elem)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_code_block(doc: Document, text: str):
    for line in text.strip("\n").splitlines():
        p = doc.add_paragraph(line if line else " ", style="Code Block")
        shade_paragraph(p, CODE_FILL)


def add_bullet(doc: Document, text: str, bold_prefix: str | None = None):
    p = doc.add_paragraph(style="List Bullet")
    if bold_prefix:
        run = p.add_run(bold_prefix)
        run.bold = True
        p.add_run(text)
    else:
        p.add_run(text)


def add_table(doc: Document, headers: list[str], rows: list[list[str]]):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        for p in hdr[i].paragraphs:
            for r in p.runs:
                r.bold = True
    for r_idx, row in enumerate(rows):
        cells = table.rows[r_idx + 1].cells
        for c_idx, val in enumerate(row):
            cells[c_idx].text = val
    doc.add_paragraph()


def pdf_pages_to_png(pdf_path: Path, out_dir: Path, dpi: int = 160) -> list[Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    paths: list[Path] = []
    pdf = fitz.open(pdf_path)
    for i, page in enumerate(pdf):
        pix = page.get_pixmap(dpi=dpi)
        path = out_dir / f"architecture_pdf_page_{i + 1}.png"
        pix.save(str(path))
        paths.append(path)
    pdf.close()
    return paths


def draw_layer_stack_diagram(path: Path) -> Path:
    """Generate EDTA layer-stack diagram as PNG."""
    width, height = 1400, 900
    img = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(img)
    try:
        font_title = ImageFont.truetype("arial.ttf", 22)
        font_box = ImageFont.truetype("arial.ttf", 14)
        font_small = ImageFont.truetype("arial.ttf", 12)
    except OSError:
        font_title = ImageFont.load_default()
        font_box = font_title
        font_small = font_title

    draw.text((40, 24), "EDTA Layer Stack — Governed Personalization Pipeline", fill="#1F4D78", font=font_title)

    layers = [
        ("Input", "#EAF6EF", "#1F7A52", ["Scenario NLP", "Profile lookup (CDP/CRM)"]),
        ("Memory & Graph", "#EDF7FB", "#21759F", ["TKGE — Temporal Knowledge Graph", "EML — Experience Memory Layer"]),
        ("Orchestration", "#FFF7E7", "#BD7B13", ["HAOE — Rules → SLM → ML → LLM"]),
        ("Scoring", "#F4F0FF", "#6C55A3", ["EDS", "Semantic + Channel fit", "TAPL", "OSE", "Final ranker"]),
        ("Travel vertical", "#FFF0F0", "#B84A4A", ["Empathy Engine — hidden needs, TCO, profiles"]),
        ("Output", "#F3F5F3", "#6B746C", ["Explanation router", "Auditable JSON + demo UI"]),
    ]

    x, box_w, box_h, gap = 80, width - 160, 95, 18
    y = 80
    for title, fill, stroke, lines in layers:
        draw.rounded_rectangle([x, y, x + box_w, y + box_h], radius=12, fill=fill, outline=stroke, width=2)
        draw.text((x + 16, y + 12), title, fill=stroke, font=font_box)
        ly = y + 38
        for line in lines:
            draw.text((x + 24, ly), f"• {line}", fill="#333333", font=font_small)
            ly += 18
        if y > 80:
            cx = x + box_w // 2
            draw.line([(cx, y - gap), (cx, y)], fill="#53645A", width=2)
            draw.polygon([(cx, y), (cx - 6, y - 10), (cx + 6, y - 10)], fill="#53645A")
        y += box_h + gap

    draw.text((40, height - 36), f"Live demo: {LIVE_DEMO}", fill="#555555", font=font_small)
    img.save(path)
    return path


def draw_four_tier_diagram(path: Path) -> Path:
    width, height = 1200, 420
    img = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(img)
    try:
        font = ImageFont.truetype("arial.ttf", 14)
        font_title = ImageFont.truetype("arial.ttf", 20)
    except OSError:
        font = ImageFont.load_default()
        font_title = font

    draw.text((40, 20), "HAOE — Four-Tier Inference (cheapest reliable tier wins)", fill="#1F4D78", font=font_title)
    tiers = [
        ("Rules", "Deterministic YAML + EDS", "#EAF6EF", "#1F7A52"),
        ("SLM", "Distilled memory + small model", "#EDF7FB", "#21759F"),
        ("ML", "sklearn intent/journey/TAPL/ranker", "#F4F0FF", "#6C55A3"),
        ("LLM", "Optional teacher (budget-gated)", "#FFF0F0", "#B84A4A"),
    ]
    x, w, h, gap = 60, 240, 120, 30
    y = 100
    for i, (name, desc, fill, stroke) in enumerate(tiers):
        bx = x + i * (w + gap)
        draw.rounded_rectangle([bx, y, bx + w, y + h], radius=10, fill=fill, outline=stroke, width=2)
        draw.text((bx + 16, y + 20), name, fill=stroke, font=font)
        draw.text((bx + 16, y + 50), desc, fill="#333333", font=font)
        if i < len(tiers) - 1:
            ax = bx + w + 4
            draw.line([(ax, y + h // 2), (ax + gap - 8, y + h // 2)], fill="#53645A", width=2)
            draw.polygon([(ax + gap - 8, y + h // 2), (ax + gap - 18, y + h // 2 - 8), (ax + gap - 18, y + h // 2 + 8)], fill="#53645A")
    draw.text((60, 280), "Unified InferenceResult: tier, confidence, rules_fired, cost_units, fallback metadata", fill="#555555", font=font)
    img.save(path)
    return path


def add_image(doc: Document, path: Path, width_in: float = 6.5, caption: str | None = None):
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width_in))
    if caption:
        doc.add_paragraph(caption, style="Caption")


def build_document():
    ASSETS.mkdir(parents=True, exist_ok=True)

    if not PDF_DIAGRAM.exists():
        import sys

        sys.path.insert(0, str(ROOT))
        import generate_architecture_pdf as gap

        gap.main()

    pdf_images = pdf_pages_to_png(PDF_DIAGRAM, ASSETS)
    layer_png = draw_layer_stack_diagram(ASSETS / "layer_stack.png")
    tier_png = draw_four_tier_diagram(ASSETS / "four_tier.png")

    doc = Document()
    configure_styles(doc)

    # Cover
    doc.add_paragraph("EDTA Architecture Document", style="Title")
    doc.add_paragraph(
        "Experience-Driven Targeting Architecture for Governed, Explainable, Trust-Aware Personalization",
        style="Doc Subtitle",
    )
    doc.add_paragraph("Author: Jerald Selvaraj", style="Metadata")
    doc.add_paragraph("Version: 1.0 — July 2026", style="Metadata")

    p = doc.add_paragraph(style="Metadata")
    p.add_run("Live demo: ")
    add_hyperlink(p, LIVE_DEMO, LIVE_DEMO)
    p = doc.add_paragraph(style="Metadata")
    p.add_run("GitHub repository: ")
    add_hyperlink(p, GITHUB, GITHUB)
    p = doc.add_paragraph(style="Metadata")
    p.add_run("API documentation: ")
    add_hyperlink(p, SWAGGER, SWAGGER)

    doc.add_page_break()

    # TOC placeholder
    doc.add_paragraph("Table of Contents", style="Heading 1")
    toc_items = [
        "1. Executive Summary",
        "2. Problem Statement",
        "3. System Context & Diagrams",
        "4. Architecture Layer Stack",
        "5. End-to-End Request Flow",
        "6. Layer-by-Layer Design",
        "7. Empathy Engine (Travel)",
        "8. Four-Tier Inference (HAOE)",
        "9. Why EDTA Is Unique",
        "10. Benefits",
        "11. Live Demo Walkthrough",
        "12. GitHub Repository Guide",
        "13. Benchmarks & Evidence",
        "14. Scope & Limits",
    ]
    for item in toc_items:
        doc.add_paragraph(item, style="List Number")
    doc.add_page_break()

    # 1 Executive Summary
    doc.add_paragraph("1. Executive Summary", style="Heading 1")
    doc.add_paragraph(
        "EDTA (Experience-Driven Targeting Architecture) is an open reference implementation that treats "
        "enterprise personalization as a governed decision pipeline—not a single black-box model. "
        "It combines temporal journey context, experience memory, four-tier AI orchestration, "
        "trust-aware policy (TAPL), outcome simulation, and optional empathy-aware travel ranking."
    )
    for acronym, rest in [
        ("TKGE", " — Temporal Knowledge Graph for live journey context"),
        ("EML", " — Experience Memory for trust, fatigue, and preferences"),
        ("HAOE", " — Hybrid AI Orchestration: Rules → SLM → ML → optional LLM"),
        ("EDS", " — Experience DNA Score for explainable relevance"),
        ("TAPL", " — Trust-Aware Personalization Layer affecting rank, not only logs"),
        ("OSE", " — Outcome Simulation Engine for conversion and revenue impact"),
        ("Empathy Engine", " — Hidden needs, route enrichment, TCO (travel demo)"),
    ]:
        add_bullet(doc, rest, acronym)
    doc.add_paragraph(
        "The LLM enriches and explains; it does not own consent, compliance, fatigue, or final trust governance."
    )

    # 2 Problem
    doc.add_paragraph("2. Problem Statement", style="Heading 1")
    add_code_block(doc, "context → model → ranked item")
    doc.add_paragraph("This pattern fails when teams need explainability, cost control, and trust enforcement:")
    add_table(
        doc,
        ["Gap", "Consequence"],
        [
            ["Opaque ranking", "Cannot answer why an offer was shown"],
            ["LLM-first design", "Unbounded cost, latency, provider lock-in"],
            ["Trust as metadata", "Consent/fatigue logged but not enforced in ranking"],
            ["Single channel", "Web, chatbot, IoT need unified policy"],
            ["No outcome view", "Optimizes clicks, not expected business outcome"],
        ],
    )

    # 3 Diagrams
    doc.add_paragraph("3. System Context & Architecture Diagrams", style="Heading 1")
    doc.add_paragraph(
        "The following diagrams illustrate runtime flow and codebase layering. "
        "Publish this document alongside the GitHub repository and live demo URL above."
    )
    add_image(doc, pdf_images[0], 6.8, "Figure 1 — End-to-end runtime flow (Browser → API → Engine → AI → Response)")
    if len(pdf_images) > 1:
        add_image(doc, pdf_images[1], 6.8, "Figure 2 — Layered codebase view (folders and responsibilities)")
    add_image(doc, layer_png, 6.5, "Figure 3 — EDTA logical layer stack")
    add_image(doc, tier_png, 6.2, "Figure 4 — HAOE four-tier inference routing")

    # 4 Layer stack text
    doc.add_paragraph("4. Architecture Layer Stack", style="Heading 1")
    add_table(
        doc,
        ["Acronym", "Name", "Role"],
        [
            ["TKGE", "Temporal Knowledge Graph Engine", "Journey timeline, inferred intent, per-scenario subjects"],
            ["EML", "Experience Memory Layer", "Trust, fatigue, preferences, outcome history"],
            ["HAOE", "Hybrid AI Orchestration Engine", "Rules / SLM / ML / LLM tier selection"],
            ["EDS", "Experience DNA Score", "Explainable relevance before ranker fusion"],
            ["TAPL", "Trust-Aware Personalization Layer", "show / soften / delay / suppress / generic_fallback"],
            ["OSE", "Outcome Simulation Engine", "Conversion, revenue, expected outcome score"],
        ],
    )

    # 5 Flow
    doc.add_paragraph("5. End-to-End Request Flow", style="Heading 1")
    doc.add_paragraph("Structured API (POST /recommend)", style="Heading 2")
    add_code_block(
        doc,
        """CustomerContext → Profile enrich → EML enrich → TKGE graph
  → HAOE tier select → For each candidate:
      EDS → Semantic → TAPL → OSE → Ranker → Empathy boost
  → Record EML → Return auditable JSON""",
    )
    doc.add_paragraph("Free-text demo (POST /recommend-from-scenario)", style="Heading 2")
    add_code_block(
        doc,
        """scenario_text → ScenarioNLPParser → scenario_anonymous_id()
  → EmpathyEngine.process() → merge vehicle catalog
  → RecommendationEngine → empathy panels + architecture summary""",
    )

    # 6 Layers
    doc.add_paragraph("6. Layer-by-Layer Design", style="Heading 1")
    layers = [
        ("6.1 Scenario NLP", "app/scenario_nlp.py", "Parses free text to CustomerContext: channel, intent, journey, consent, events."),
        ("6.2 Profile Lookup", "app/profile_service.py", "Merges CRM/CDP data (cust-789 demo) outside the ranker."),
        ("6.3 TKGE", "app/context_graph.py", "Builds temporal graph; live_journey_sequence; per-profile anonymous IDs."),
        ("6.4 EML", "app/eml_store.py", "SQLite trust/fatigue/preferences; OSE calibration from historical outcomes."),
        ("6.5 EDS", "app/eds.py", "Intent, engagement, business value, journey, context overlap, risk adjustment."),
        ("6.6 TAPL", "app/ai/tapl_model.py", "Policy YAML overrides; rank caps for suppress/delay; SQLite audit."),
        ("6.7 OSE", "app/ai/outcome_model.py", "Conversion, revenue, trust/journey impact; heuristic+ML blend."),
        ("6.8 Ranker", "app/recommender.py", "Hybrid score + empathy boost; ai_rank_score vs final_hybrid_score."),
        ("6.9 Explanation", "app/inference/explanation_router.py", "SLM first, LLM escalation when enabled."),
    ]
    for title, path, desc in layers:
        doc.add_paragraph(title, style="Heading 2")
        p = doc.add_paragraph()
        r = p.add_run(f"Implementation: {path} — ")
        r.font.name = "Consolas"
        r.font.size = Pt(9)
        p.add_run(desc)

    # 7 Empathy
    doc.add_paragraph("7. Empathy Engine (Travel Vertical)", style="Heading 1")
    doc.add_paragraph(
        "Eleven trained scenarios (cust-789 loyalty rental + 10 road-trip archetypes) are defined in "
        "config/empathy/trained_scenarios.yaml and config/empathy/scenario_profiles.yaml."
    )
    add_table(
        doc,
        ["#", "Profile", "Expected vehicle", "Objective"],
        [
            ["0", "loyalty_family_rental", "family_friendly_suv", "cust-789 SFO family SUV rental"],
            ["1", "winter_mountain_denver", "awd_suv", "Winter Denver 500-mile trip"],
            ["2", "long_family_vacation", "large_family_suv", "NJ → Orlando family of five"],
            ["3", "business_executive", "luxury_sedan", "35k annual business miles"],
            ["4", "national_park_adventure", "awd_suv", "Multi-park adventure"],
            ["5", "urban_commuter", "hybrid_midsize", "18k urban commute"],
            ["6", "electric_vehicle", "electric_midsize", "First EV purchase"],
            ["7", "luxury_winter_suv", "premium_suv", "Colorado luxury winter"],
            ["8", "first_time_driver", "economy_compact", "New driver"],
            ["9", "rental_seattle_vacation", "awd_suv", "Seattle 7-day rental"],
            ["10", "wet_weather_hurricane", "awd_suv", "Hurricane season Southeast"],
        ],
    )

    # 8 HAOE
    doc.add_paragraph("8. Four-Tier Inference (HAOE)", style="Heading 1")
    doc.add_paragraph(
        "HAOE (app/orchestration.py) selects the cheapest tier that meets confidence and policy thresholds. "
        "Session cost budget and LLM circuit breaker prevent runaway token spend."
    )

    # 9 Unique
    doc.add_paragraph("9. Why EDTA Is Unique", style="Heading 1")
    add_table(
        doc,
        ["Typical pattern", "EDTA difference"],
        [
            ["LLM-first personalization", "LLM optional; TAPL + rules own governance"],
            ["Monolithic ranker", "EDS + TAPL + OSE + ranker — each auditable"],
            ["Trust as logging", "TAPL modifies rank and action"],
            ["Hidden model routing", "Public tiers with InferenceResult contract"],
            ["Stateless API", "EML + TKGE across sessions"],
            ["Vertical rewrite", "Same core; swap YAML rules, catalog, empathy profiles"],
        ],
    )

    # 10 Benefits
    doc.add_paragraph("10. Benefits", style="Heading 1")
    doc.add_paragraph("Engineering", style="Heading 2")
    for b in [
        "149+ pytest cases; travel benchmark with 100% vehicle alignment",
        "Modular adapters for profile, rules, catalog",
        "Graceful degradation: LLM off → ML/rules → heuristics",
        "OpenAPI + stable request_summary contract",
    ]:
        add_bullet(doc, b)
    doc.add_paragraph("Product & compliance", style="Heading 2")
    for b in [
        "Explain why this offer, why now, why this channel",
        "Demonstrate consent and fatigue (TAPL delay/suppress)",
        "Healthcare scope: HCP education routing — not clinical decision support",
    ]:
        add_bullet(doc, b)

    # 11 Demo
    doc.add_paragraph("11. Live Demo Walkthrough", style="Heading 1")
    p = doc.add_paragraph()
    p.add_run("URL: ")
    add_hyperlink(p, LIVE_DEMO, LIVE_DEMO)
    doc.add_paragraph("5-minute demo script:", style="Heading 3")
    for step in [
        "Default cust-789 scenario — loyalty family SUV, high trust, loyalty_family_rental profile",
        "Paste Winter Denver scenario — awd_suv, mountain constraints",
        "Add 'personalization consent false' — TAPL generic_fallback",
        "Add 'fatigue count 8' — TAPL delay",
        "Scroll to benchmark table — 11 scenarios with distinct scores",
    ]:
        add_bullet(doc, step)

    # 12 GitHub
    doc.add_paragraph("12. GitHub Repository Guide", style="Heading 1")
    p = doc.add_paragraph()
    add_hyperlink(p, GITHUB, GITHUB)
    add_table(
        doc,
        ["Path", "Purpose"],
        [
            ["app/recommender.py", "Main ranking loop"],
            ["app/services/recommendation_handlers.py", "API handlers + scenario route"],
            ["app/empathy/", "Empathy engine + trained profiles"],
            ["app/context_graph.py", "TKGE"],
            ["app/eml_store.py", "EML persistence"],
            ["config/empathy/", "Scenario profiles and trained corpus"],
            ["static/scenario.html", "Browser demo UI"],
            ["tests/", "pytest suite"],
            ["scripts/benchmark_travel_scenarios.py", "Regenerate score matrix"],
        ],
    )
    add_code_block(
        doc,
        """pip install -r requirements.txt
uvicorn app.main:app --reload
# Demo: http://localhost:8000/scenario-demo
python scripts/benchmark_travel_scenarios.py
pytest""",
    )

    # 13 Benchmarks
    doc.add_paragraph("13. Benchmarks & Evidence", style="Heading 1")
    doc.add_paragraph(
        "Travel scenario matrix: docs/TRAVEL_SCENARIO_SCORES.md — 11 scenarios, 100% vehicle match. "
        "Four-tier benchmarks: docs/BENCHMARKS.md. API: GET /travel-scenario-benchmark."
    )

    # 14 Limits
    doc.add_paragraph("14. Scope & Honest Limits", style="Heading 1")
    for lim in [
        "Reference implementation on demonstration training data—not Netflix-scale production traffic.",
        "ML models use heuristic blend when predictions saturate.",
        "Healthcare: approved HCP education routing only; not clinical decision support.",
        "Hosted demo latency includes Render cold-start overhead.",
    ]:
        add_bullet(doc, lim)

    # Footer
    section = doc.sections[0]
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run(f"EDTA Architecture — {GITHUB} — {LIVE_DEMO}")
    run.font.size = Pt(8)
    run.font.color.rgb = MUTED

    doc.save(OUTPUT)
    print(OUTPUT.resolve())


if __name__ == "__main__":
    build_document()
